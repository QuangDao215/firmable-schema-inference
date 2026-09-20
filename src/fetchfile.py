"""Download a resource and look inside it.

Used by the hand-check in Part 1 and by stage C1 and C2 in Part 2. The point is
to answer "what is this file really" from its bytes, not from the format label
a publisher typed into CKAN.

Nothing here calls a model.
"""

from __future__ import annotations

import csv
import io
import json
import re
import time
import zipfile
from pathlib import Path

import requests

USER_AGENT = "firmable-assignment/0.1 (schema inference take-home)"
DEFAULT_MAX_BYTES = 8_000_000          # 8 MB is plenty to see a file's shape


def download(url: str, dest: Path, max_bytes: int = DEFAULT_MAX_BYTES) -> dict:
    """Fetch up to max_bytes of a URL. Returns what we actually got.

    We stream and stop early rather than trusting the size in the catalogue,
    which is often missing or wrong.
    """
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        for attempt in range(3):
            reply = requests.get(url, stream=True, timeout=90,
                                 allow_redirects=True,
                                 headers={"User-Agent": USER_AGENT})
            # 202 means the portal is building the file on demand. Waiting and
            # asking again usually works. CKAN datastore dumps do this.
            if reply.status_code == 202 and attempt < 2:
                reply.close()
                time.sleep(5 * (attempt + 1))
                continue
            break

        with reply:
            status = reply.status_code
            content_type = reply.headers.get("Content-Type", "")
            if status != 200:
                return {"ok": False, "status": status, "error": f"HTTP {status}"}

            written = 0
            with open(dest, "wb") as f:
                for chunk in reply.iter_content(64 * 1024):
                    f.write(chunk)
                    written += len(chunk)
                    if written >= max_bytes:
                        break
    except Exception as err:
        return {"ok": False, "status": 0, "error": f"{type(err).__name__}: {err}"}

    return {"ok": True, "status": status, "content_type": content_type,
            "bytes": written, "truncated": written >= max_bytes,
            "path": str(dest)}


def sniff(path: Path) -> str:
    """What the first bytes say this file is, ignoring its name and label."""
    head = path.open("rb").read(4096)
    if head[:2] == b"PK":
        return "ZIP"                      # xlsx and docx are zips too
    if head[:5] == b"%PDF-":
        return "PDF"
    text = head.decode("utf-8", errors="replace").lstrip()
    if text[:1] in "{[":
        return "JSON"
    if text[:5].lower() == "<?xml" or text[:1] == "<":
        return "HTML" if re.match(r"<(!doctype )?html", text[:20], re.I) else "XML"
    if "\t" in text.split("\n")[0] and "," not in text.split("\n")[0]:
        return "TSV"
    return "CSV"


def _preview_delimited(path: Path, delimiter: str, rows: int) -> dict:
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        collected = []
        for i, row in enumerate(reader):
            collected.append(row)
            if i >= rows:
                break
    if not collected:
        return {"kind": "empty"}
    header = find_header_row(collected)
    return {"kind": "table", "header_row": header,
            "columns": collected[header], "rows": collected[header + 1:]}


def find_header_row(grid: list[list], look_at: int = 6) -> int:
    """Which row holds the column names.

    Spreadsheets often open with a title banner, a blank row, then the real
    headers. We take the row with the most non-empty cells among the first
    few, which is the header in every messy sheet we have seen. The answer is
    recorded in the mapping config, so a human can disagree with it.
    """
    best_row, best_filled = 0, -1
    for i, row in enumerate(grid[:look_at]):
        filled = sum(1 for cell in row if str(cell).strip())
        if filled > best_filled:
            best_row, best_filled = i, filled
    return best_row


def _preview_xlsx(path: Path, rows: int) -> dict:
    from openpyxl import load_workbook
    # Pass a file object, not a path. openpyxl refuses a path whose extension
    # it does not recognise, and our cache files are named .bin.
    book = load_workbook(path.open("rb"), read_only=True, data_only=True)
    sheet_names = list(book.sheetnames)
    sheet = book[sheet_names[0]]
    collected = []
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        collected.append(["" if v is None else str(v) for v in row])
        if i >= rows:
            break
    sheet_title = sheet.title
    book.close()
    if not collected:
        return {"kind": "empty"}
    header = find_header_row(collected)
    # Trailing empty cells are Excel padding, not columns.
    columns = collected[header]
    while columns and not str(columns[-1]).strip():
        columns.pop()
    return {"kind": "table", "sheet": sheet_title, "sheets": sheet_names,
            "header_row": header, "columns": columns,
            "rows": collected[header + 1:]}


def _preview_json(path: Path, rows: int) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        # A JSON Lines file, or a truncated download.
        records = []
        for line in text.splitlines()[:rows]:
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                break
        if records:
            return {"kind": "records", "records": records,
                    "columns": sorted({k for r in records if isinstance(r, dict)
                                       for k in r})}
        return {"kind": "unreadable_json"}

    if isinstance(data, list):
        records = data[:rows]
    elif isinstance(data, dict):
        # Find the longest list of objects anywhere near the top.
        best = max((v for v in data.values() if isinstance(v, list)),
                   key=len, default=None)
        records = (best or [data])[:rows]
    else:
        records = [data]
    return {"kind": "records", "records": records,
            "columns": sorted({k for r in records if isinstance(r, dict) for k in r})}


def _preview_xml(path: Path, rows: int) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    tags = re.findall(r"<([A-Za-z_][\w.-]*)[ >]", text)
    common = {}
    for tag in tags:
        common[tag] = common.get(tag, 0) + 1
    return {"kind": "xml",
            "top_tags": sorted(common.items(), key=lambda kv: -kv[1])[:15],
            "head": text[:2500]}


def _preview_zip(path: Path, rows: int) -> dict:
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        return {"kind": "truncated_zip",
                "note": "download was cut short, cannot open the archive"}

    names = archive.namelist()
    if any(n.startswith("xl/") for n in names):
        archive.close()
        return {**_preview_xlsx(path, rows), "really": "XLSX"}

    inner = next((n for n in names
                  if n.lower().endswith((".csv", ".tsv", ".txt", ".json", ".xml"))),
                 None)
    if not inner:
        archive.close()
        return {"kind": "zip", "files": names[:25]}

    raw = archive.open(inner).read(400_000)
    archive.close()
    text = raw.decode("utf-8", errors="replace")
    if inner.lower().endswith(".xml"):
        return {"kind": "xml", "inner_file": inner, "head": text[:2500],
                "files_in_zip": names[:25]}
    if inner.lower().endswith(".json"):
        return {"kind": "zip_json", "inner_file": inner, "head": text[:2500],
                "files_in_zip": names[:25]}
    reader = csv.reader(io.StringIO(text),
                        delimiter="\t" if inner.endswith(".tsv") else ",")
    collected = [row for _, row in zip(range(rows + 1), reader)]
    if not collected:
        return {"kind": "empty", "inner_file": inner}
    return {"kind": "table", "inner_file": inner, "files_in_zip": names[:25],
            "columns": collected[0], "rows": collected[1:]}


ZIP_MAX_BYTES = 80_000_000        # a zip must be whole to be opened


def download_for_preview(url: str, dest: Path,
                         max_bytes: int = DEFAULT_MAX_BYTES) -> dict:
    """Download enough of a file to look inside it.

    A zip keeps its index at the end, so a partial download cannot be opened
    at all. xlsx files are zips. So we fetch a sample, and if the first bytes
    say zip, we go back for the whole file.
    """
    got = download(url, dest, max_bytes)
    if not got.get("ok"):
        return got
    if sniff(dest) == "ZIP" and got.get("truncated"):
        got = download(url, dest, ZIP_MAX_BYTES)
        got["refetched_whole_zip"] = True
    return got


def preview(path: Path, rows: int = 8) -> dict:
    """Read the first few records, whatever kind of file this is."""
    kind = sniff(path)
    try:
        if kind == "ZIP":
            view = _preview_zip(path, rows)
            # An xlsx is a zip. Say xlsx, because that is what a reader needs
            # to know and what the mapping config has to record.
            return {**view, "sniffed": view.pop("really", kind)}
        if kind == "JSON":
            return {**_preview_json(path, rows), "sniffed": kind}
        if kind == "XML":
            return {**_preview_xml(path, rows), "sniffed": kind}
        if kind == "TSV":
            return {**_preview_delimited(path, "\t", rows), "sniffed": kind}
        if kind in ("PDF", "HTML"):
            return {"kind": kind.lower(), "sniffed": kind,
                    "note": "not a data file"}
        return {**_preview_delimited(path, ",", rows), "sniffed": kind}
    except Exception as err:
        return {"kind": "error", "sniffed": kind,
                "error": f"{type(err).__name__}: {err}"}


# --------------------------------------------------------------------------
# Deterministic signals: is there a business in here?
# --------------------------------------------------------------------------

COMPANY_WORDS = re.compile(
    r"\b(pty\.? ?ltd|proprietary limited|limited|ltd\.?|incorporated|inc\.?|"
    r"holdings|group|enterprises|services|trading as|t/a|partnership|"
    r"co-?operative|association)\b", re.I)
ABN_SHAPED = re.compile(r"\b\d{11}\b")
ACN_SHAPED = re.compile(r"\b\d{9}\b")


def business_signals(view: dict) -> dict:
    """Count things in real data that only appear when businesses are named.

    This is evidence, not a verdict. A file full of 'Pty Ltd' almost certainly
    names companies. A file with none might still, under a different naming
    convention.
    """
    if view.get("kind") == "table":
        text = " | ".join(" ".join(str(c) for c in row)
                          for row in view.get("rows", []))
        columns = view.get("columns", [])
    elif view.get("kind") == "records":
        text = json.dumps(view.get("records", []), default=str)
        columns = view.get("columns", [])
    else:
        text = view.get("head", "")
        columns = [t for t, _ in view.get("top_tags", [])]

    column_text = " ".join(str(c) for c in columns).lower()
    return {
        "company_word_hits": len(COMPANY_WORDS.findall(text)),
        "abn_shaped_values": len(ABN_SHAPED.findall(text)),
        "acn_shaped_values": len(ACN_SHAPED.findall(text)),
        "name_like_columns": [c for c in columns if re.search(
            r"name|entity|company|business|organisation|trading|licensee|"
            r"supplier|contractor|abn|acn", str(c), re.I)],
        "column_count": len(columns),
        "columns_mention_business": bool(re.search(
            r"abn|acn|company|business|entity|licensee|supplier|contractor",
            column_text)),
    }


# --------------------------------------------------------------------------
# Reading every record out of a file, once we know what it is
# --------------------------------------------------------------------------

def read_records(path: Path, shape: dict, limit: int = 2000) -> list[dict]:
    """Pull records out of a file using what the probe worked out about it.

    `shape` is the resource section of a mapping config: real_format, encoding,
    delimiter, header_row, sheet, zip_member and so on. The engine uses the
    same function, so what the agent profiled is exactly what gets extracted.
    """
    fmt = shape.get("real_format", "CSV")

    if fmt in ("CSV", "TSV"):
        return _records_delimited(path, shape, limit)
    if fmt == "XLSX":
        return _records_xlsx(path, shape, limit)
    if fmt in ("JSON", "JSONL", "GEOJSON"):
        return _records_json(path, shape, limit)
    if fmt == "XML":
        return _records_xml(path, shape, limit)
    if fmt == "ZIP":
        return _records_zip(path, shape, limit)
    raise ValueError(f"cannot read records from format {fmt!r}")


def _rows_to_records(rows, header_row: int, limit: int) -> list[dict]:
    header, records = None, []
    for i, row in enumerate(rows):
        if i < header_row:
            continue
        cells = ["" if c is None else str(c) for c in row]
        if header is None:
            while cells and not cells[-1].strip():
                cells.pop()
            header = [c.strip().lstrip("\ufeff") or f"column_{n}"
                      for n, c in enumerate(cells)]
            continue
        records.append({name: (cells[n] if n < len(cells) else "")
                        for n, name in enumerate(header)})
        if len(records) >= limit:
            break
    return records


def _records_delimited(path: Path, shape: dict, limit: int) -> list[dict]:
    delimiter = shape.get("delimiter") or ("\t" if shape.get("real_format") == "TSV" else ",")
    with open(path, newline="", encoding=shape.get("encoding", "utf-8"),
              errors="replace") as f:
        return _rows_to_records(csv.reader(f, delimiter=delimiter),
                                shape.get("header_row", 0), limit)


def _records_xlsx(path: Path, shape: dict, limit: int) -> list[dict]:
    from openpyxl import load_workbook
    book = load_workbook(path.open("rb"), read_only=True, data_only=True)
    sheet = book[shape["sheet"]] if shape.get("sheet") in book.sheetnames \
        else book[book.sheetnames[0]]
    records = _rows_to_records(sheet.iter_rows(values_only=True),
                               shape.get("header_row", 0), limit)
    book.close()
    return records


def _dig(data, path: str):
    for piece in (path or "").split("."):
        if not piece:
            continue
        data = data.get(piece, {}) if isinstance(data, dict) else {}
    return data


def _records_json(path: Path, shape: dict, limit: int) -> list[dict]:
    text = path.read_text(encoding=shape.get("encoding", "utf-8"), errors="replace")
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        records = []
        for line in text.splitlines():
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                continue
            if len(records) >= limit:
                break
        return records

    if shape.get("json_record_path"):
        data = _dig(data, shape["json_record_path"])
    elif isinstance(data, dict):
        data = max((v for v in data.values() if isinstance(v, list)),
                   key=len, default=[data])
    rows = data if isinstance(data, list) else [data]
    return [r for r in rows[:limit] if isinstance(r, dict)]


def _records_xml(path: Path, shape: dict, limit: int) -> list[dict]:
    import xml.etree.ElementTree as ET
    tag = shape.get("xml_record_tag")
    records = []
    for _, element in ET.iterparse(path, events=("end",)):
        name = element.tag.split("}")[-1]
        if tag and name != tag:
            continue
        if not tag and len(element) == 0:
            continue
        row = {**{k.split("}")[-1]: v for k, v in element.attrib.items()}}
        for child in element:
            key = child.tag.split("}")[-1]
            row[key] = (child.text or "").strip()
            for k, v in child.attrib.items():
                row[f"{key}@{k.split('}')[-1]}"] = v
        if row:
            records.append(row)
        element.clear()
        if len(records) >= limit:
            break
    return records


def _records_zip(path: Path, shape: dict, limit: int) -> list[dict]:
    archive = zipfile.ZipFile(path)
    names = archive.namelist()
    if any(n.startswith("xl/") for n in names):
        archive.close()
        return _records_xlsx(path, {**shape, "real_format": "XLSX"}, limit)

    member = shape.get("zip_member") or next(
        (n for n in names if n.lower().endswith((".csv", ".tsv", ".txt", ".json", ".xml"))),
        None)
    if not member:
        archive.close()
        raise ValueError(f"no readable file inside the archive: {names[:8]}")

    import tempfile
    with archive.open(member) as inner:
        with tempfile.NamedTemporaryFile(delete=False) as tmp:
            tmp.write(inner.read(60_000_000))
            tmp_path = Path(tmp.name)
    archive.close()
    inner_format = ("JSON" if member.lower().endswith(".json") else
                    "XML" if member.lower().endswith(".xml") else
                    "TSV" if member.lower().endswith(".tsv") else "CSV")
    return read_records(tmp_path, {**shape, "real_format": inner_format}, limit)
